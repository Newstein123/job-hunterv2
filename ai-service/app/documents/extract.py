"""Plain-text extraction from existing resumes (PDF, DOCX, TXT, MD)."""

import re
from pathlib import Path

from app.models.resume import Resume, ResumeLayout

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedDocumentError(ValueError):
    pass


# A CV counts as multi-column when over a quarter of its sentence-length text starts right of this point.
_RIGHT_COLUMN_X = 0.38
_MULTI_COLUMN_SHARE = 0.25
_CONTACT = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+|\+\s?\d{1,3}[\s\d()-]{6,}")


def _extract_pdf(path: Path) -> tuple[str, int, ResumeLayout]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages, right, total = [], 0, 0
    for page in reader.pages:
        width = float(page.mediabox.width) or 1.0
        runs: list[tuple[float, int]] = []

        def visit(text, cm, tm, _font, _size, runs=runs):
            if len(text.strip()) >= 25:  # sentence-length runs; short ones (dates, labels) don't define columns
                runs.append((cm[4] + tm[4] * cm[0], len(text.strip())))

        pages.append(page.extract_text(visitor_text=visit) or "")
        for x, n in runs:
            total += n
            right += n if x > _RIGHT_COLUMN_X * width else 0
    multi = total >= 300 and right / total > _MULTI_COLUMN_SHARE
    return "\n".join(pages), len(reader.pages), ResumeLayout(multi_column=multi)


def _extract_docx(path: Path) -> tuple[str, ResumeLayout]:
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    tables = 0
    for table in document.tables:
        if len(table.columns) > 1:
            tables += 1
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    # Text boxes aren't in document.paragraphs; Word stores each one twice (modern + legacy), so dedupe.
    boxes, seen = 0, set(parts)
    for box in document.element.body.iter(qn("w:txbxContent")):
        text = "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in box.iter(qn("w:p"))).strip()
        if text and text not in seen:
            seen.add(text)
            parts.append(text)
            boxes += 1

    body = "\n".join(parts)
    header = "\n".join(p.text for section in document.sections
                       for part in (section.header, section.footer) for p in part.paragraphs)
    layout = ResumeLayout(
        tables=tables, text_boxes=boxes,
        contact_only_in_header=bool(_CONTACT.search(header)) and not _CONTACT.search(body),
    )
    return body, layout


def extract_text(path: Path) -> tuple[str, int | None]:
    """Return (text, page_count). page_count is only known for PDFs."""
    text, pages, _ = extract(path)
    return text, pages


def extract(path: Path) -> tuple[str, int | None, ResumeLayout | None]:
    """Text, page count (PDF only) and layout features."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        text, layout = _extract_docx(path)
        return text, None, layout
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8"), None, None
    raise UnsupportedDocumentError(f"Unsupported resume format: {path.name}")


def load_resume(path: Path) -> Resume:
    text, pages, layout = extract(path)
    if not text.strip():
        raise UnsupportedDocumentError(f"No extractable text in {path.name} (scanned PDF?)")
    return Resume(
        id=path.stem,
        filename=path.name,
        format=path.suffix.lower().lstrip("."),
        text=text,
        page_count=pages,
        layout=layout,
    )


def load_resume_bytes(resume_id: str, filename: str, content: bytes) -> Resume:
    """Load a resume sent over the API. The caller's id is kept so results map back."""
    import tempfile

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocumentError(f"Unsupported resume format: {filename}")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / f"resume{suffix}"
        path.write_bytes(content)
        resume = load_resume(path)
    return resume.model_copy(update={"id": resume_id, "filename": filename})


def load_resumes(directory: Path) -> list[Resume]:
    paths = sorted(p for p in directory.iterdir() if p.suffix.lower() in SUPPORTED_SUFFIXES)
    if not paths:
        raise FileNotFoundError(f"No resumes found in {directory}")
    return [load_resume(p) for p in paths]

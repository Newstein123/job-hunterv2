from app.llm.client import FakeLLMClient
from app.models.resume import Resume, ResumeLayout
from app.pipelines.resume_review import review_resume, run_checks
from app.rules.models import CVRules

# A fictional CV in the course layout.
CV = """AYE THANDAR
Quantity Surveyor
+95 9 400 000 001 (WhatsApp) | aye@example.org

CAREER SUMMARY
After graduating with a Bachelor of Engineering (Civil), I have worked as a Quantity Surveyor for 6 years.

WORK EXPERIENCE
Quantity Surveyor, Golden Build Co., Ltd 2021 - Present
Duties and Responsibilities: Prepared Bills of Quantities for residential towers
Assistant Quantity Surveyor, Shwe Pyi Engineering 2018 - 2021
Checked tender documents for completeness

EDUCATION
Technological University, Bachelor of Engineering 2012 - 2017

SKILLS
AutoCAD, CostX, Microsoft Excel

PERSONAL INFO
Date of Birth - 1.1.1995  Nationality - Myanmar  Passport No - MX000000
Expected Salary - 2500 SGD (Nego)   Availability - Immediately
LANGUAGE
English (Intermediate), Myanmar (Native)
I hereby declare that the above information is true.
""" + "Detail line. " * 20


def _resume(text=CV, filename="Aye_Thandar_Quantity_Surveyor.pdf", fmt="pdf", pages=2):
    return Resume(id="1", filename=filename, format=fmt, text=text, page_count=pages)


def test_course_checklist_passes_a_complete_cv():
    checks = {c.id: c for c in run_checks(_resume(), CVRules(), "Aye Thandar")}
    assert {k for k, c in checks.items() if c.status != "pass"} == set(), {k: c.detail for k, c in checks.items() if c.status != "pass"}
    assert "WhatsApp" in checks["whatsapp"].evidence
    assert checks["dates"].evidence.startswith("2021–Present")  # education years are not job dates


def test_checklist_flags_what_is_missing_and_bad_file_names():
    bare = "JOHN DOE\njohn@example.org\nEXPERIENCE\nClerk 2015 - 2016\nManager 2019 - 2022\n" + "word " * 80
    checks = {c.id: c for c in run_checks(_resume(bare, "Resume - Copy (1).docx", "docx", None), CVRules(), "John Doe")}
    for missing in ("phone", "whatsapp", "personal_info", "passport", "salary", "availability", "languages", "summary", "declaration"):
        assert checks[missing].status == "missing", missing
    assert checks["dates"].status == "check"  # oldest job listed first
    assert checks["format"].status == "check"
    assert checks["file_name"].status == "check" and "John_Doe_Job_Title.pdf" in checks["file_name"].detail


def test_scanned_cv_is_flagged_and_no_ai_call_is_made():
    llm = FakeLLMClient([])
    review = review_resume(_resume("scan", pages=1), CVRules(), llm, "Aye Thandar")
    assert [c.id for c in review.checks] == ["readable", "file_name"] and review.checks[0].status == "missing"
    assert review.suggestions == [] and llm.calls == []


def test_suggestions_must_quote_the_cv_and_examples_cannot_add_numbers():
    llm = FakeLLMClient([{"suggestions": [
        {"title": "Lead with results", "detail": "Say what your BQs achieved.",
         "quote": "Duties and Responsibilities: Prepared Bills of Quantities", "example": "Prepared Bills of Quantities for 3 towers"},
        {"title": "Invented", "detail": "x", "quote": "Managed a team of 40 engineers", "example": ""},
        {"title": "Add a headline", "detail": "Put your target role under your name.", "quote": "", "example": ""},
    ]}])
    review = review_resume(_resume(), CVRules(), llm, "Aye Thandar")
    assert [s.title for s in review.suggestions] == ["Lead with results", "Add a headline"]
    assert review.suggestions[0].example is None  # "3" is not in the CV
    assert any("Managed a team" in w for w in review.warnings)
    assert llm.calls[0]["think"] is False and "CHECKLIST_ALREADY_SHOWN" in llm.calls[0]["prompt"]


def test_examples_that_add_claims_are_dropped():
    llm = FakeLLMClient([{"suggestions": [
        {"title": "Say more", "detail": "d", "quote": "Checked tender documents for completeness",
         "example": "Checked tender documents for completeness and accuracy, ensuring compliance with project specifications"},
    ]}])
    review = review_resume(_resume(), CVRules(), llm, "Aye Thandar")
    assert review.suggestions[0].example is None
    assert "never suggest removing" in llm.calls[0]["system"]


def test_ats_checks_flag_columns_boxes_header_contact_and_headings():
    two_col = _resume().model_copy(update={"layout": ResumeLayout(multi_column=True)})
    checks = {c.id: c for c in run_checks(two_col, CVRules(), "Aye Thandar")}
    assert checks["ats_columns"].status == "check" and checks["ats_columns"].group == "ats"
    assert checks["ats_headings"].status == "pass"

    no_skills = _resume(CV.replace("SKILLS", "EXPERTIZE"), "Aye_Thandar_QS.docx", "docx", None).model_copy(
        update={"layout": ResumeLayout(text_boxes=3, contact_only_in_header=True)})
    checks = {c.id: c for c in run_checks(no_skills, CVRules(), "Aye Thandar")}
    assert checks["ats_headings"].status == "check" and "Skills" in checks["ats_headings"].detail
    assert checks["ats_boxes"].status == "check" and "3 text box" in checks["ats_boxes"].detail
    assert checks["ats_header_contact"].status == "check"
    assert "ats_columns" not in checks  # column detection is for PDFs


def test_docx_text_boxes_are_read_and_counted(tmp_path):
    import copy
    import docx
    from docx.oxml.ns import qn
    from lxml import etree
    from app.documents.extract import load_resume

    path = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Quantity Surveyor")
    run = document.add_paragraph().add_run()
    box = etree.SubElement(run._r, qn("w:pict"))
    content = etree.SubElement(etree.SubElement(box, qn("w:txbxContent")), qn("w:p"))
    etree.SubElement(etree.SubElement(content, qn("w:r")), qn("w:t")).text = "Prepared Bills of Quantities"
    document.save(str(path))

    resume = load_resume(path)
    assert "Prepared Bills of Quantities" in resume.text
    assert resume.layout.text_boxes == 1

import pytest

from app.documents.extract import UnsupportedDocumentError, load_resume
from app.skills.normalize import find_skills_in_text, normalize_skill, normalize_skills, text_mentions


def test_aliases_normalize_to_canonical():
    assert normalize_skill("Postgres") == "PostgreSQL"
    assert normalize_skill("RESTful APIs") == "REST API"
    assert normalize_skill("k8s") == "Kubernetes"
    assert normalize_skill("Some Niche Tool") == "Some Niche Tool"


def test_normalize_skills_dedupes():
    assert normalize_skills(["JS", "javascript", "React", "reactjs"]) == ["JavaScript", "React"]


def test_text_matching_respects_boundaries():
    assert text_mentions("We use Node.js and C# daily", "Node.js")
    assert text_mentions("We use Node.js and C# daily", "C#")
    assert not text_mentions("JavaScript developer", "Java")
    assert not text_mentions("the rest of the team", "REST API")
    assert not text_mentions("ready to go live", "Go")
    assert text_mentions("Services written in Go", "Go")
    assert find_skills_in_text("Laravel + MySQL, deployed with docker", ["mysql", "Docker", "AWS"]) == [
        "MySQL",
        "Docker",
    ]


def test_extract_docx(tmp_path):
    import docx

    path = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Laravel developer")
    document.save(str(path))
    resume = load_resume(path)
    assert resume.format == "docx" and "Laravel" in resume.text


def test_extract_pdf(tmp_path):
    fpdf = pytest.importorskip("fpdf")
    path = tmp_path / "cv.pdf"
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "PHP and Docker")
    pdf.output(str(path))
    resume = load_resume(path)
    assert resume.page_count == 1 and "Docker" in resume.text


def test_unsupported_format(tmp_path):
    path = tmp_path / "cv.odt"
    path.write_text("x")
    with pytest.raises(UnsupportedDocumentError):
        load_resume(path)


def test_industry_vocabulary_from_course_resumes():
    assert normalize_skill("MEP") == "M&E" and normalize_skill("Building Services") == "M&E"
    assert normalize_skill("HVAC") == "ACMV"
    assert normalize_skill("BOQ") == "Bill of Quantities"
    assert normalize_skill("autodesk revit") == "Revit"
    assert find_skills_in_text("Prepared BQ and handled tenders using AutoCAD 2D", ["Bill of Quantities", "Tendering", "AutoCAD"]) == [
        "Bill of Quantities", "Tendering", "AutoCAD"]


def test_common_words_are_not_mistaken_for_software():
    assert not text_mentions("Ensure a safe working environment", "SAFE")
    assert text_mentions("Foundation design in SAFE and ETABS", "SAFE")
    assert not text_mentions("I will tally the results", "Tally")
    assert text_mentions("Experience with SAP and MYOB", "SAP")
